import os
import shutil
import tempfile
import unittest
import re
import zipfile
from datetime import datetime, timezone

import openpyxl

from statement_checker import analyze

HEADER = ["STT", "Ngày GD", "Nội dung", "Ghi nợ", "Ghi có", "Số dư"]
TXNS = [
    (1, datetime(2024, 1, 2), "Luong thang 12", None, 15000000, 25000000),
    (2, datetime(2024, 1, 5), "Chuyen tien nha", 4500000, None, 20500000),
    (3, datetime(2024, 1, 10), "Thanh toan dien", 750000, None, 19750000),
    (4, datetime(2024, 1, 15), "Nhan tien ban hang", None, 3250000, 23000000),
]


def build_xlsx(path, txns=TXNS, creator="CoreBanking", modified_by=None, app="CoreBanking Export",
               formula=False, hide_row=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SaoKe"
    ws.append(["SAO KÊ TÀI KHOẢN", "0123456789"])
    ws.append(["Số dư đầu kỳ", 10000000])
    ws.append(HEADER)
    for t in txns:
        ws.append(list(t))
    ws.append(["Số dư cuối kỳ", txns[-1][5]])
    for row in ws.iter_rows(min_row=4, max_row=3 + len(txns), min_col=2, max_col=2):
        row[0].number_format = "dd/mm/yyyy"
    if formula:
        ws["F6"] = "=F5-D6"
    if hide_row:
        ws.row_dimensions[hide_row].hidden = True
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = modified_by or creator
    wb.properties.created = datetime.now(timezone.utc).replace(tzinfo=None)  # openpyxl stamps modified=now on save
    wb.save(path)
    _set_app(path, app)


def _set_app(path, app):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = re.sub(rb"<Application>[^<]*</Application>",
                              f"<Application>{app}</Application>".encode(), data)
            dst.writestr(item, data)
    shutil.move(tmp, path)


def severities(report, check=None):
    return [f.severity for f in report.findings if check is None or f.check == check]


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def path(self, name):
        return os.path.join(self.tmp, name)

    def test_genuine_xlsx_is_low_risk(self):
        p = self.path("goc.xlsx")
        build_xlsx(p)
        report = analyze(p)
        self.assertEqual(report.score, 0, [f.message for f in report.findings if f.severity != "info"])
        self.assertTrue(any("khớp với toàn bộ" in f.message for f in report.findings))

    def test_edited_balance_in_xlsx(self):
        txns = list(TXNS)
        txns[2] = (3, datetime(2024, 1, 10), "Thanh toan dien", 750000, None, 91750000)
        p = self.path("sua.xlsx")
        build_xlsx(p, txns=txns)
        self.assertIn("critical", severities(analyze(p), "content"))

    def test_resaved_by_excel_with_formula_and_hidden_row(self):
        p = self.path("luu-lai.xlsx")
        build_xlsx(p, modified_by="Nguyen Van A", app="Microsoft Excel", formula=True, hide_row=5)
        report = analyze(p)
        messages = " | ".join(f.message for f in report.findings)
        self.assertIn("Người sửa cuối khác người tạo", messages)
        self.assertIn("ô công thức", messages)
        self.assertIn("dòng dữ liệu bị ẩn", messages)
        self.assertGreaterEqual(report.score, 50)

    def test_csv_reconciles(self):
        p = self.path("sao-ke.csv")
        with open(p, "w", encoding="utf-8") as fh:
            fh.write("Ngày,Nội dung,Ghi nợ,Ghi có,Số dư\n")
            fh.write("02/01/2024,Luong,,15000000,25000000\n")
            fh.write("05/01/2024,Chuyen tien,4500000,,20500000\n")
            fh.write("10/01/2024,Dien,750000,,19750000\n")
        report = analyze(p)
        self.assertTrue(any("khớp với toàn bộ" in f.message for f in report.findings))
        self.assertLess(report.score, 20)

    def test_csv_edited_balance(self):
        p = self.path("sua.csv")
        with open(p, "w", encoding="utf-8") as fh:
            fh.write("Ngày;Nội dung;Ghi nợ;Ghi có;Số dư\n")
            fh.write("02/01/2024;Luong;;15000000;25000000\n")
            fh.write("05/01/2024;Chuyen tien;4500000;;80500000\n")
            fh.write("10/01/2024;Dien;750000;;19750000\n")
        self.assertIn("critical", severities(analyze(p), "content"))


if __name__ == "__main__":
    unittest.main()
