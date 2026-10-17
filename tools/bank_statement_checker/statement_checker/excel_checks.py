"""Checks for statements exported as spreadsheets (.xlsx, .csv).

Each sheet row is turned into a text line (dates as dd/mm/yyyy, money with
thousands separators) so the same reconciliation as for PDFs can run on it.
"""

import csv
import re
import zipfile
from datetime import date, datetime, timedelta

from . import content_checks

BALANCE_HEADER = re.compile(r"số dư|so du|balance", re.IGNORECASE)
MONEY_HEADER = re.compile(
    r"số dư|so du|balance|ghi nợ|ghi có|ghi no|ghi co|nợ|có|debit|credit|số tiền|so tien|amount|"
    r"rút|nộp|gửi|phát sinh|withdraw|deposit|thu|chi", re.IGNORECASE)
DESKTOP_APPS = ["microsoft excel", "microsoft macintosh excel", "wps", "libreoffice", "openoffice",
                "numbers", "google sheets", "kingsoft"]
# Libraries that bank systems use to generate spreadsheets (some report themselves as "Excel compatible").
EXPORT_LIBRARIES = ["openpyxl", "apache poi", "npoi", "epplus", "aspose", "phpspreadsheet", "phpexcel",
                    "sheetjs", "closedxml", "xlsxwriter", "jxls", "jasperreports", "compatible"]


def _find_header(rows):
    """Return (row index, set of money column indexes) of the header row, or (None, None)."""
    for i, row in enumerate(rows[:40]):
        texts = [str(v) if isinstance(v, str) else "" for v in row]
        labels = [t for t in texts if t.strip() and len(t) < 40]
        if len(labels) >= 3 and any(BALANCE_HEADER.search(t) for t in labels):
            return i, {j for j, t in enumerate(texts) if t and len(t) < 40 and MONEY_HEADER.search(t)}
    return None, None


def _cell_text(value, money):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        # Money gets separators so it is recognised as an amount; other numbers
        # (row numbers, account numbers) are written plainly and ignored.
        return f"{value:,.2f}" if money else str(int(value)) if float(value).is_integer() else str(value)
    return str(value).strip()


def _is_date(value):
    return isinstance(value, (datetime, date)) or (
        isinstance(value, str) and bool(content_checks.DATE_RE.search(value)))


def rows_to_lines(rows):
    header, money_cols = _find_header(rows)
    lines = []
    for row in rows:
        # Inside the transaction table only the money columns count; on other rows
        # (opening/closing balance lines) any number may be an amount.
        in_table = bool(money_cols) and any(_is_date(v) for v in row)
        cells = []
        for j, value in enumerate(row):
            money = (j in money_cols) if in_table else isinstance(value, (int, float)) and abs(value) < 1e12
            cells.append(_cell_text(value, money))
        line = "  ".join(c for c in cells if c)
        if line:
            lines.append(line)
    return lines, header, money_cols


def check_xlsx(path, report):
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
        values_wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        report.add("structure", "high", f"File Excel bị lỗi, không đọc được: {exc}")
        return

    _check_xlsx_properties(path, wb, report)

    lines, formula_cells, hidden_rows, hidden_cols, text_money = [], [], [], [], []
    hidden_sheets = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        vws = values_wb[ws.title]
        rows = [list(r) for r in vws.iter_rows(values_only=True)]
        sheet_lines, header, money_cols = rows_to_lines(rows)
        lines.extend(sheet_lines)

        for r_idx, row in enumerate(ws.iter_rows(), 1):
            has_date = any(_is_date(c.value) for c in row)
            for c_idx, cell in enumerate(row):
                if has_date and cell.data_type == "f":
                    formula_cells.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
                if (has_date and money_cols and c_idx in money_cols and isinstance(cell.value, str)
                        and content_checks.AMOUNT_RE.fullmatch(cell.value.strip())):
                    text_money.append(f"{ws.title}!{cell.coordinate}: '{cell.value}'")
            if ws.row_dimensions[r_idx].hidden and any(c.value is not None for c in row):
                hidden_rows.append(f"{ws.title}!dòng {r_idx}")
        hidden_cols += [f"{ws.title}!cột {k}" for k, d in ws.column_dimensions.items() if d.hidden]

    if formula_cells:
        report.add("excel", "medium",
                   f"Có {len(formula_cells)} ô công thức trong các dòng giao dịch - sao kê xuất từ ngân hàng "
                   "chứa giá trị, không chứa công thức tự tính", formula_cells[:5])
    if hidden_rows:
        report.add("excel", "high", f"Có {len(hidden_rows)} dòng dữ liệu bị ẩn - có thể dùng để giấu giao dịch",
                   hidden_rows[:5])
    if hidden_cols:
        report.add("excel", "low", f"Có {len(hidden_cols)} cột bị ẩn", hidden_cols[:5])
    if hidden_sheets:
        report.add("excel", "low", f"Có sheet bị ẩn: {', '.join(hidden_sheets)}")
    if text_money:
        report.add("excel", "medium",
                   f"Có {len(text_money)} ô số tiền lưu dạng chữ thay vì số - dấu hiệu gõ tay vào file",
                   text_money[:5])

    content_checks.check_transactions(lines, report)


def _check_xlsx_properties(path, wb, report):
    props = wb.properties
    app = ""
    try:
        with zipfile.ZipFile(path) as z:
            m = re.search(rb"<Application>([^<]*)</Application>", z.read("docProps/app.xml"))
            app = m.group(1).decode("utf-8", "replace") if m else ""
    except (KeyError, zipfile.BadZipFile):
        pass

    report.add("metadata", "info", "Thông tin file Excel", [
        f"Ứng dụng: {app or '(trống)'}",
        f"Người tạo: {props.creator or '(trống)'}",
        f"Người sửa cuối: {props.lastModifiedBy or '(trống)'}",
        f"Ngày tạo: {props.created or '(trống)'}",
        f"Ngày sửa: {props.modified or '(trống)'}",
    ])
    if props.lastModifiedBy and props.creator and props.lastModifiedBy != props.creator:
        report.add("metadata", "medium", "Người sửa cuối khác người tạo file - file đã được mở và lưu lại",
                   [f"Tạo: {props.creator}; sửa cuối: {props.lastModifiedBy}"])
    elif props.lastModifiedBy and not props.creator:
        report.add("metadata", "low", f"File có người sửa cuối ({props.lastModifiedBy}) nhưng trống người tạo")
    if props.created and props.modified and props.modified - props.created > timedelta(minutes=5):
        report.add("metadata", "medium", "Ngày sửa khác ngày tạo - file đã được lưu lại sau khi xuất",
                   [f"Chênh lệch: {props.modified - props.created}"])
    if any(a in app.lower() for a in DESKTOP_APPS) and not any(l in app.lower() for l in EXPORT_LIBRARIES):
        report.add("metadata", "low",
                   f"File được lưu bằng phần mềm bảng tính ({app}) - bản xuất từ hệ thống ngân hàng thường do thư viện tạo ra")


def check_csv(path, report):
    with open(path, "rb") as fh:
        raw = fh.read()
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[_number_or_text(v) for v in row] for row in csv.reader(text.splitlines(), dialect)]
    report.add("metadata", "low",
               "File CSV không có metadata và rất dễ sửa bằng tay - nên yêu cầu bản PDF gốc từ ngân hàng")
    lines, _, _ = rows_to_lines(rows)
    content_checks.check_transactions(lines, report)


def _number_or_text(value):
    v = value.strip()
    if re.fullmatch(r"[-+]?\d+(\.\d+)?", v) and len(v) < 16:
        return float(v)
    return v
